import re
from datetime import datetime, date
import pandas as pd
from django.utils.dateparse import parse_date
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Required column aliases
COLUMN_ALIASES = {
    'email': ['email', 'e-mail', 'mail', 'البريد'],
    'password': ['password', 'pass', 'كلمة السر', 'كلمة المرور'],
    'passport_number': ['passport_number', 'passport', 'passport_no', 'رقم الجواز'],
    'phone_number': ['phone_number', 'phone', 'mobile', 'رقم الهاتف', 'الموبايل'],
    'birth_date': ['birth_date', 'dob', 'date_of_birth', 'birthdate', 'تاريخ الميلاد'],
    'nationality': ['nationality', 'country', 'الجنسية'],
    'full_name': ['full_name', 'name', 'customer_name', 'الاسم'],
    'passport_image_path': ['passport_image_path', 'image_path', 'passport_image', 'صورة الجواز']
}

EMAIL_REGEX = re.compile(r'^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$')

def normalize_column_name(col_name: str) -> str:
    cleaned = str(col_name).strip().lower().replace(' ', '_')
    for canonical, aliases in COLUMN_ALIASES.items():
        if cleaned in aliases:
            return canonical
    return cleaned

def parse_flexible_date(val):
    if val is None or pd.isna(val):
        return None
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    val_str = str(val).strip()
    # Try Django parse_date (YYYY-MM-DD)
    parsed = parse_date(val_str)
    if parsed:
        return parsed
    # Try standard formats
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d', '%m/%d/%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(val_str, fmt).date()
        except ValueError:
            continue
    return None

def validate_and_preview_excel(file_obj):
    filename = getattr(file_obj, 'name', 'upload.xlsx').lower()
    try:
        if filename.endswith('.csv'):
            df = pd.read_csv(file_obj)
        else:
            df = pd.read_excel(file_obj)
    except Exception as e:
        return {
            'success': False,
            'error': f'Failed to parse file: {str(e)}',
            'total_rows': 0,
            'valid_count': 0,
            'invalid_count': 0,
            'valid_rows': [],
            'invalid_rows': []
        }

    # Normalize column names
    col_mapping = {col: normalize_column_name(col) for col in df.columns}
    df = df.rename(columns=col_mapping)

    required = ['email', 'password', 'passport_number', 'phone_number', 'birth_date', 'nationality']
    missing = [c for c in required if c not in df.columns]
    if missing:
        return {
            'success': False,
            'error': f'Missing required columns: {", ".join(missing)}',
            'missing_columns': missing,
            'total_rows': len(df),
            'valid_count': 0,
            'invalid_count': len(df),
            'valid_rows': [],
            'invalid_rows': []
        }

    valid_rows = []
    invalid_rows = []

    for idx, row in df.iterrows():
        excel_row_num = idx + 2  # 1-based, skipping header
        row_errors = []

        # Check if entire row is empty
        row_values = [str(val).strip() for val in row.values if pd.notna(val) and str(val).strip() != '']
        if not row_values:
            continue

        # Validate Email
        email_val = str(row.get('email', '')).strip() if pd.notna(row.get('email')) else ''
        if not email_val:
            row_errors.append('Missing email address')
        elif not EMAIL_REGEX.match(email_val):
            row_errors.append(f'Invalid email format: {email_val}')

        # Validate Password
        password_val = str(row.get('password', '')).strip() if pd.notna(row.get('password')) else ''
        if not password_val:
            row_errors.append('Missing password')

        # Validate Passport Number
        passport_val = str(row.get('passport_number', '')).strip().upper() if pd.notna(row.get('passport_number')) else ''
        if not passport_val:
            row_errors.append('Missing passport number')
        elif len(passport_val) < 5 or len(passport_val) > 20:
            row_errors.append(f'Invalid passport number length ({len(passport_val)} chars): {passport_val}')

        # Validate Nationality
        nationality_val = str(row.get('nationality', '')).strip() if pd.notna(row.get('nationality')) else ''
        if not nationality_val:
            row_errors.append('Missing nationality')

        # Validate Phone Number
        phone_val = str(row.get('phone_number', '')).strip() if pd.notna(row.get('phone_number')) else ''
        if not phone_val:
            row_errors.append('Missing phone number')

        # Validate Birth Date
        raw_dob = row.get('birth_date')
        parsed_dob = parse_flexible_date(raw_dob)
        if not parsed_dob:
            row_errors.append(f'Invalid birth date format: "{raw_dob}". Expected YYYY-MM-DD or DD/MM/YYYY')

        # Optional full name
        name_val = str(row.get('full_name', '')).strip() if pd.notna(row.get('full_name')) else ''
        image_path = str(row.get('passport_image_path', '')).strip() if pd.notna(row.get('passport_image_path')) else ''

        customer_summary = {
            'row_number': excel_row_num,
            'full_name': name_val or email_val.split('@')[0],
            'email': email_val,
            'passport_number': passport_val,
            'nationality': nationality_val,
            'phone_number': phone_val,
            'birth_date': parsed_dob.strftime('%Y-%m-%d') if parsed_dob else str(raw_dob),
            'passport_image_path': image_path or None,
        }

        if row_errors:
            customer_summary['errors'] = row_errors
            invalid_rows.append(customer_summary)
        else:
            # Keep raw password internally for confirmation import
            customer_summary['password'] = password_val
            valid_rows.append(customer_summary)

    return {
        'success': True,
        'total_rows': len(valid_rows) + len(invalid_rows),
        'valid_count': len(valid_rows),
        'invalid_count': len(invalid_rows),
        'valid_rows': valid_rows,
        'invalid_rows': invalid_rows
    }

def generate_sample_excel_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customers"

    headers = [
        'full_name', 'email', 'password', 'passport_number',
        'phone_number', 'birth_date', 'nationality', 'passport_image_path'
    ]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sample_rows = [
        ["Ahmed Mohamed", "ahmed.user1@gmail.com", "Password123!", "A12345678", "+201012345678", "1992-05-15", "Egypt", ""],
        ["Sara Ali", "sara.user2@gmail.com", "SecurePass456!", "A87654321", "+201123456789", "1995-10-20", "Egypt", ""],
        ["Omar Khaled", "omar.user3@gmail.com", "StrongKey789!", "A11223344", "+201234567890", "1988-03-12", "Egypt", ""],
    ]
    for r in sample_rows:
        ws.append(r)

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
